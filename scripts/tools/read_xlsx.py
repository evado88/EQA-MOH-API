"""Prints an .xlsx sheet by sheet, for the PT forms."""
import openpyxl
p = r"C:\Users\nkole\Downloads\CDL TB Scheme Form and tools\CDL TB Scheme Form and tools\CDL-PT-F-008 Xpert PT Results form.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
for ws in wb.worksheets:
    print("=== SHEET:", ws.title, ws.dimensions)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,80)):
        vals = [f"{c.coordinate}={c.value!r}" for c in row if c.value is not None]
        if vals:
            print(" | ".join(vals))
    try:
        for dv in ws.data_validations.dataValidation:
            print("DV:", dv.sqref, dv.type, dv.formula1)
    except Exception as e:
        print("dv err", e)
